from pathlib import Path
import json

from pptx import Presentation


pptx_path = Path("FodFin AGPR.pptx")
prs = Presentation(str(pptx_path))
slides = []

for idx, slide in enumerate(prs.slides, start=1):
    notes_text = ""
    try:
        text_runs = []
        notes_slide = slide.notes_slide
        for shape in notes_slide.shapes:
            if not hasattr(shape, "text_frame") or shape.text_frame is None:
                continue
            text = shape.text_frame.text.strip()
            if text:
                text_runs.append(text)
        notes_text = "\n".join(text_runs).strip()
    except Exception as exc:
        notes_text = f"__ERROR__ {exc}"

    slides.append(
        {
            "slide": idx,
            "notes": notes_text,
        }
    )

out = Path("slide_notes.json")
out.write_text(json.dumps(slides, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"slides={len(slides)}")
print(out.read_text(encoding="utf-8"))

# Made with Bob
